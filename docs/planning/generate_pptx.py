import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_customer_proposal():
    prs = Presentation()
    # 16:9 ワイド画面設定
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # カラーパレット（洗練されたコンサル・ハイエンド提案書仕様）
    C_NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A
    C_NAVY_PRIMARY = RGBColor(30, 58, 138)  # #1E3A8A
    C_GOLD = RGBColor(202, 138, 4)          # #CA8A04
    C_ACCENT_BLUE = RGBColor(37, 99, 235)   # #2563EB
    C_BG_LIGHT = RGBColor(248, 250, 252)    # #F8FAFC
    C_BORDER = RGBColor(226, 232, 240)      # #E2E8F0
    C_TEXT_MAIN = RGBColor(15, 23, 42)      # #0F172A
    C_TEXT_MUTED = RGBColor(71, 85, 105)    # #475569
    C_WHITE = RGBColor(255, 255, 255)       # #FFFFFF
    C_CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF

    def add_header(slide, title_text, category="STRATEGIC E-COMMERCE PROPOSAL"):
        # 上部アクセントバー
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_GOLD
        top_bar.line.fill.background()

        # ヘッダーテキスト
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = C_GOLD

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY_DARK

    def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # ==========================================
    # SLIDE 1: 表紙（正式な顧客提示用）
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY_DARK
    bg1.line.fill.background()

    # アクセント装飾
    dec1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.8), Inches(0.15), Inches(3.2))
    dec1.fill.solid()
    dec1.fill.fore_color.rgb = C_GOLD
    dec1.line.fill.background()

    tbox1 = s1.shapes.add_textbox(Inches(1.4), Inches(1.8), Inches(10.5), Inches(3.2))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "PROPOSAL FOR DIGITAL BUSINESS EXPANSION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_GOLD
    
    p = tf1.add_paragraph()
    p.text = "創業百二十年の伝統と先進デジタル技術の融合\nオムニチャネルECプラットフォーム構築 企画提案書"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p = tf1.add_paragraph()
    p.text = "ご提案先: 有限会社 権田原水産　代表取締役 権田原 権三 様"
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(226, 232, 240)

    info_box = s1.shapes.add_textbox(Inches(1.4), Inches(5.8), Inches(10), Inches(1.0))
    itf = info_box.text_frame
    ip = itf.paragraphs[0]
    ip.text = "提出日: 2026年9月吉日  |  株式会社ハイパーネクストソリューションズ DXソリューション事業部  |  Ver 1.0 (PROPOSAL)"
    ip.font.size = Pt(11)
    ip.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # SLIDE 2: プロジェクトサマリー & 導入ビジョン
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "プロジェクト背景および戦略的導入ビジョン")

    col_w = Inches(3.64)
    gap = Inches(0.4)
    start_x = Inches(0.8)
    top_y = Inches(1.5)
    card_h = Inches(5.2)

    cards_data = [
        ("01 / 市場環境と事業機会", "老舗ブランドの全国D2C展開", [
            "・築地・浅草で培った120年の歴史と顧客信頼をアセット化",
            "・中元・歳暮需要から日常の家庭食・贈答需要への拡張",
            "・地理的制約を打破し、全国の食通層へのダイレクトアクセス確立",
            "・伝統の味を次世代へ継承するブランドストーリーのデジタル発信",
            "・既存店舗事業との相乗効果による企業価値の最大化"
        ]),
        ("02 / 目指すべき主要成果 (KPI)", "持続可能なオムニチャネル成長", [
            "・初年度オンライン売上高の大幅伸長と収益基盤の複線化",
            "・デジタル接点を通じた若年・共働き世帯の新規ファン獲得",
            "・ブランド体験の深化によるリピート購入率の最大化",
            "・直販化（D2C）による高収益体質へのシフトと中間マージン削減",
            "・全国的なブランドプレゼンスの確立とメディア波及効果"
        ]),
        ("03 / 推進方針と提供価値", "スピードと品質を両立する導入", [
            "・モダンなWeb技術（Laravel / Docker）による高拡張性基盤",
            "・顧客心理に寄り添ったUI/UXによるシームレスな購買体験",
            "・店頭とEC双方のシナジーを追求した柔軟な業務設計",
            "・最短リードタイムでの市場投入を実現する段階的アプローチ",
            "・将来の外部モール連携・API接続を見据えた拡張設計"
        ])
    ]

    for i, (tag, title, points) in enumerate(cards_data):
        cx = start_x + i * (col_w + gap)
        add_card(s2, cx, top_y, col_w, card_h, bg_color=RGBColor(248, 250, 252))
        tb = s2.shapes.add_textbox(cx + Inches(0.25), top_y + Inches(0.2), col_w - Inches(0.5), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_GOLD
        
        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_DARK

        for pt in points:
            p = tf.add_paragraph()
            p.text = pt
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_TEXT_MAIN

    # ==========================================
    # SLIDE 3: クライアントヒアリング議事録（企画側の拡大解釈・曖昧さゼロの表現）
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "顧客ヒアリング要件の構造化と戦略的再定義")

    grid_w = Inches(5.66)
    grid_h = Inches(2.4)
    
    hearings = [
        ("【FV映像戦略】ブランドアイデンティティの直感認知", 
         "【店主様ご要望】「秘伝のタレが煮立つ臨場感を動画と音響で画面一杯に伝えたい」\n"
         "【弊社戦略解釈】ファーストビューに全画面高精細シズル動画を採用。視覚・聴覚に直接訴求するリッチメディア体験により、来訪と同時に120年の歴史と職人の情熱へ没入させる「エモーショナル・イマーシブ演出」として具現化します。"),
        
        ("【ロイヤルティ戦略】VIPコミュニティとゲーミフィケーション", 
         "【店主様ご要望】「昔からの馴染み客を優遇し、特別な繋がりを維持したい」\n"
         "【弊社戦略解釈】合言葉入力による特別オファー開放施策を導入。顧客の自発的エンゲージメントを喚起するゲーミフィケーション要素として設計し、ロイヤル顧客の特別感を醸成、熱狂的ファン層のLTV最大化を図ります。"),

        ("【品質管理戦略】老舗ならではのハイタッチ・コンシェルジュ", 
         "【店主様ご要望】「できたての無添加品質を最優先で、確実に美味しい状態でお届けしたい」\n"
         "【弊社戦略解釈】鮮度管理が極めてシビアな生鮮佃煮について、機械的な自動出荷ではなく個別直接確認を伴うコンシェルジュ型配送オペレーションを採用。安心感と特別感を演出する高付加価値接客と位置付けます。"),

        ("【決済・データ戦略】フリクションレス決済 & 会員基盤構築", 
         "【店主様ご要望】「幅広い客層の購買に対応し、将来のリピート基盤を整備したい」\n"
         "【弊社戦略解釈】次世代キャッシュレス決済（QR・カード）を段階的に導入。購買行動データを蓄積する会員ポイントプログラムを整備し、将来的なオムニチャネルCRM展開の礎を築きます。")
    ]

    positions = [
        (Inches(0.8), Inches(1.5)),
        (Inches(6.86), Inches(1.5)),
        (Inches(0.8), Inches(4.2)),
        (Inches(6.86), Inches(4.2))
    ]

    for (title, body), (gx, gy) in zip(hearings, positions):
        add_card(s3, gx, gy, grid_w, grid_h, bg_color=C_WHITE, border_color=C_BORDER)
        tb = s3.shapes.add_textbox(gx + Inches(0.25), gy + Inches(0.2), grid_w - Inches(0.5), grid_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_PRIMARY

        p = tf.add_paragraph()
        p.text = body
        p.font.size = Pt(10)
        p.font.color.rgb = C_TEXT_MAIN

    # ==========================================
    # SLIDE 4: システム要件定義書（まともな設計＋ビジネス懸念点）
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "システム要件定義および設計仕様概要（現行設計方針）")

    req_table = [
        ("在庫管理基盤", "排他制御付きリアルタイム在庫管理", "トランザクション分離レベルを最適化し、同時注文時の過剰受注（マイナス在庫）を完全に防止する行ロック機構を設計。", "※店頭販売との在庫連携につきましては、POS未導入期間は管理画面での手動差異照合運用を推奨いたします。"),
        ("注文・カート", "厳格な型検証と整合性バリデーション", "購入数量の正数検証（1〜99個制限）、負数・特殊文字の自動遮断、金額改ざんを防止するサービストランザクション。", "※大口注文（法人様向けギフト等）の需要については、別途個別問い合わせ窓口の併設を推奨いたします。"),
        ("税額計算処理", "インボイス制度完全準拠の端数処理", "関係法令および適格請求書保存方式に基づき、明細ごとではなく一伝票ごとに税額を算出・統一する端数管理基盤。", "※貴社既存の基幹会計システムまたは顧問税理士様の端数処理方針（切捨て等）との最終すり合わせを推奨いたします。"),
        ("セキュリティ", "OWASP Top 10準拠の堅牢設計", "Prepared Statementによる生SQLインジェクション完全遮断、Blade標準エスケープによるXSS防御、厳格な認可制御(RBAC)。", "※管理者アカウントの適切な運用管理（複雑なパスワードの設定および定時変更）の遵守をお願い申し上げます。"),
        ("ギフト対応", "熨斗・贈答オプションの柔軟な入力設計", "慶事・弔事・名入れ等に対応する備考フォームの設置および、手書き熨斗業務を支援する管理画面印刷用出力。", "※繁忙期（中元・歳暮期）の名入れ記載漏れを低減するため、主要用途のプリセット選択肢導入を検討可能です。")
    ]

    rows = 6
    cols = 4
    t_left = Inches(0.8)
    t_top = Inches(1.5)
    t_w = Inches(11.733)
    t_h = Inches(5.2)

    table_shape = s4.shapes.add_table(rows, cols, t_left, t_top, t_w, t_h)
    table = table_shape.table
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(3.2)
    table.columns[2].width = Inches(3.6)
    table.columns[3].width = Inches(2.933)

    headers = ["機能領域", "設計仕様（要件定義）", "システム実現方針", "ご検討・運用留意事項"]
    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

    for ri, (c1, c2, c3, c4) in enumerate(req_table):
        for ci, val in enumerate([c1, c2, c3, c4]):
            cell = table.cell(ri + 1, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_WHITE if ri % 2 == 0 else C_BG_LIGHT
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(9.5)
            p.font.color.rgb = C_TEXT_MAIN

    # ==========================================
    # SLIDE 5: サイトマップ設計（ビジネス文書として整然とした導線設計）
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "ECサイト 情報アーキテクチャ (IA) およびサイトマップ設計")

    sm_nodes = [
        ("TOPページ\n(ブランドFV動画 / 新着情報 / おすすめ商品)", Inches(4.8), Inches(1.5), Inches(3.7), Inches(0.85), C_NAVY_PRIMARY, C_WHITE),
        
        ("商品一覧・検索\n(/products)", Inches(0.8), Inches(2.7), Inches(2.6), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("創業の歴史・こだわり\n(/kodawari)", Inches(3.8), Inches(2.7), Inches(2.6), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("店舗案内・アクセス\n(/access)", Inches(6.8), Inches(2.7), Inches(2.6), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("会員マイページ\n(/mypage)", Inches(9.8), Inches(2.7), Inches(2.6), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),

        ("商品詳細・仕様\n(/products/{id})", Inches(0.8), Inches(3.9), Inches(2.6), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("買い物かご\n(/cart)", Inches(3.8), Inches(3.9), Inches(2.6), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("ご注文・決済手続き\n(/checkout)", Inches(6.8), Inches(3.9), Inches(2.6), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("ご注文完了・サンクス\n(/thanks)", Inches(9.8), Inches(3.9), Inches(2.6), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),

        ("【特別企画】シークレットVIPラウンジ\n(合言葉認証済みロイヤル顧客限定コンテンツ /secret-lounge)", Inches(0.8), Inches(5.1), Inches(5.6), Inches(1.2), RGBColor(254, 243, 199), C_NAVY_DARK),
        ("【コンプライアンス基盤】特定商取引法に基づく表記 / プライバシーポリシー / 利用規約\n(全ページフッター共通導線整備)", Inches(6.8), Inches(5.1), Inches(5.6), Inches(1.2), C_BG_LIGHT, C_NAVY_DARK)
    ]

    for title, x, y, w, h, bg, tc in sm_nodes:
        shape = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg
        shape.line.color.rgb = C_GOLD if bg == RGBColor(254, 243, 199) else C_BORDER
        shape.line.width = Pt(1.5) if bg == RGBColor(254, 243, 199) else Pt(1)
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = title
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = tc

    # ==========================================
    # SLIDE 6: ローンチスケジュール（絶妙に守れそうで守れない高効率アジャイル工程）
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "ローンチ推進スケジュール（最短市場投入アジャイル工程）")

    sched_data = [
        ("Phase 1: 要件確定・設計FIX", "Week 1", [
            "・事業要件および画面ワイヤー確定",
            "・DB設計・商品マスタフォーマットFIX",
            "※貴社仕様検収期間: 1営業日設定",
            "※素材（画像・原稿）の即日支給前提"
        ]),
        ("Phase 2: 実装・API連携", "Week 2", [
            "・Laravelフレームワーク基盤構築",
            "・フロントUI実装・決済連携モック接続",
            "※並行開発によるリードタイム極小化",
            "※仕様追加・変更は原則凍結"
        ]),
        ("Phase 3: 結合検証・受入検収", "Week 3", [
            "・シナリオテスト・決済疎通確認",
            "・本番商品マスタデータ登録",
            "※貴社受入テスト(UAT)期間: 2営業日",
            "※軽微な微調整の反映完了"
        ]),
        ("Phase 4: 本番移行・OPEN", "Week 4", [
            "・DNS切り替え・SSL証明書適用",
            "・本番決済稼働確認・大安吉日公開",
            "※運用開始および初期監視",
            "※初期改善イテレーションへ移行"
        ])
    ]

    sw = Inches(2.65)
    sgap = Inches(0.35)
    for i, (phase, period, points) in enumerate(sched_data):
        sx = Inches(0.8) + i * (sw + sgap)
        sy = Inches(1.6)
        sh = Inches(4.3)

        add_card(s6, sx, sy, sw, sh, bg_color=C_WHITE)
        tb = s6.shapes.add_textbox(sx + Inches(0.2), sy + Inches(0.2), sw - Inches(0.4), sh - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"STEP {i+1}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_GOLD

        p = tf.add_paragraph()
        p.text = phase
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_DARK

        p = tf.add_paragraph()
        p.text = period
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_PRIMARY

        for pt in points:
            p = tf.add_paragraph()
            p.text = pt
            p.font.size = Pt(9.5)
            p.font.color.rgb = C_TEXT_MAIN

    # 下部注記（ビジネス文章で絶妙にリスクを顧客側にヘッジする文言）
    note_box = s6.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.9))
    ntf = note_box.text_frame
    ntf.word_wrap = True
    np = ntf.paragraphs[0]
    np.text = "【推進上の前提事項およびリスクヘッジ条項】"
    np.font.size = Pt(9.5)
    np.font.bold = True
    np.font.color.rgb = C_TEXT_MUTED

    np2 = ntf.add_paragraph()
    np2.text = "※本スケジュールは最短での市場投入を達成するため、各工程における貴社ご判断および素材提供を「即日」頂戴することを前提とした高効率アジャイル工程となっております。各社審査（クレジットカード等決済代行会社）の所要日数により、一部決済手段の稼働開始がリリース以降順次となる場合がございます。"
    np2.font.size = Pt(9)
    np2.font.color.rgb = C_TEXT_MUTED

    # 保存
    output_path = "/Users/gsgwr/.gemini/antigravity/scratch/gondawara-tsukudani/docs/planning/gondawara_dx_proposal.pptx"
    prs.save(output_path)
    print(f"Customer proposal saved successfully to {output_path}")

if __name__ == "__main__":
    create_customer_proposal()
