import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 ワイド画面設定
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # カラーパレット定義
    C_NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A
    C_NAVY_PRIMARY = RGBColor(30, 58, 138)  # #1E3A8A
    C_ORANGE = RGBColor(234, 88, 12)        # #EA580C
    C_BG_LIGHT = RGBColor(248, 250, 252)    # #F8FAFC
    C_BORDER = RGBColor(226, 232, 240)      # #E2E8F0
    C_TEXT_MAIN = RGBColor(15, 23, 42)      # #0F172A
    C_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
    C_WHITE = RGBColor(255, 255, 255)       # #FFFFFF
    C_CARD_BG = RGBColor(241, 245, 249)     # #F1F5F9
    C_ALERT_BG = RGBColor(254, 242, 242)    # #FEF2F2
    C_ALERT_BORDER = RGBColor(252, 165, 165)# #FCA5A5

    def add_header(slide, title_text, category="GONDAWARA TSUKUDANI DX PROJECT"):
        # 上部アクセントバー
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_ORANGE
        top_bar.line.fill.background()

        # ヘッダーテキスト
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = C_ORANGE

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY_DARK

    def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # ==========================================
    # SLIDE 1: 表紙
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY_DARK
    bg1.line.fill.background()

    # アクセント装飾
    dec1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.8), Inches(0.15), Inches(3.2))
    dec1.fill.solid()
    dec1.fill.fore_color.rgb = C_ORANGE
    dec1.line.fill.background()

    tbox1 = s1.shapes.add_textbox(Inches(1.4), Inches(1.8), Inches(10.5), Inches(3.2))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "PROPOSAL FOR DIGITAL TRANSFORMATION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_ORANGE
    
    p = tf1.add_paragraph()
    p.text = "伝統とデジタルの融合による\n全国販路拡大 ECプラットフォーム構築計画書"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p = tf1.add_paragraph()
    p.text = "有限会社 権田原水産 / 創業明治38年「権田原佃煮店」公式オンラインショップ"
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(148, 163, 184)

    info_box = s1.shapes.add_textbox(Inches(1.4), Inches(5.8), Inches(10), Inches(1.0))
    itf = info_box.text_frame
    ip = itf.paragraphs[0]
    ip.text = "日付: 2026年9月吉日  |  作成: 株式会社ハイパーネクストソリューションズ (格安Web制作チーム)  |  Ver 1.0 (FIX)"
    ip.font.size = Pt(11)
    ip.font.color.rgb = RGBColor(203, 213, 225)

    # ==========================================
    # SLIDE 2: プロジェクトサマリー & 背景
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "プロジェクト背景 & エグゼクティブサマリー")

    # 3カラムカード
    col_w = Inches(3.64)
    gap = Inches(0.4)
    start_x = Inches(0.8)
    top_y = Inches(1.5)
    card_h = Inches(5.2)

    cards_data = [
        ("01 / 背景と課題", "老舗の危機感とDXの胎動", [
            "・創業120年、築地・浅草の店頭販売中心で推移",
            "・近隣の同業他社がネット通販を開始し売上伸長",
            "・店主権三氏（62歳）が一念発起『うちもやるぞ！』",
            "・『一番安く、すぐ作れる会社』として弊社に特命発注",
            "・IT専門部署・担当者は不在（店主が全権指示）"
        ]),
        ("02 / 掲げられた目標 (KPI)", "驚異的な急成長シナリオ", [
            "・初年度売上目標：前年比 1,000% 増（根拠レス）",
            "・『全国の若者に佃煮ブームを巻き起こす』",
            "・SNSでバズらせて毎分100件の注文殺到を想定",
            "・リピート率：80%以上を根拠なく試算",
            "・予算：業界相場の約1/10（超突貫格安プライス）"
        ]),
        ("03 / 基本方針とリスク", "スピード最優先・仕様は後追い", [
            "・『細かいことは作りながら考える』アジャイル（？）",
            "・既存業務フロー（手書き台帳）との連携は未検証",
            "・店主の居酒屋トークをそのまま要件定義に採用",
            "・品質保証（QA）工程は予算都合によりカット",
            "・来週月曜日の大安に無理やりローンチを予定"
        ])
    ]

    for i, (tag, title, points) in enumerate(cards_data):
        cx = start_x + i * (col_w + gap)
        add_card(s2, cx, top_y, col_w, card_h)
        tb = s2.shapes.add_textbox(cx + Inches(0.2), top_y + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_ORANGE
        
        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_DARK

        for pt in points:
            p = tf.add_paragraph()
            p.text = pt
            p.font.size = Pt(11)
            p.font.color.rgb = C_TEXT_MAIN

    # ==========================================
    # SLIDE 3: クライアントヒアリング議事録（ツッコミ満載）
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "顧客ヒアリング議事録（店主・権三氏の生の声）")

    # 4つのツッコミカード（2x2グリッド）
    grid_w = Inches(5.66)
    grid_h = Inches(2.4)
    
    hearings = [
        ("【トップ演出】動画バーン！音ドーン！", 
         "「トップページを開いたらな、うちの秘伝のタレがぐつぐつ煮立ってる動画を画面いっぱいに流してくれ！音もドーンと鳴らして迫力満点にしたいんだよ。今の若い奴は動画だろ？」",
         "※ブラウザの自動再生制限やスマホ通信量への配慮は未定。店主の強いこだわり。"),
        
        ("【常連割引】いい感じの合言葉で半額", 
         "「昔からの馴染み客には特別扱いしたいんだ。『権三の友達』とか合言葉を入れたら半額になるようにしといて。クーポン？よく分からんからいい感じに作っといてよ。」",
         "※合言葉がSNSに流出した場合の原価割れリスク・1回限り制限などは検討外。"),

        ("【配送エリア】生モノは電話で断る", 
         "「うちの佃煮は無添加だから足が早い。北海道や沖縄は届く前に傷んじゃうかもな。頼めないようにするか、まあ注文入ったらこっちから電話して断るからいいや！」",
         "※システム自動除外か運用手動対応か不明。電話不通時の返金仕様も存在せず。"),

        ("【決済・会員】PayPayとか全部＆ポイント適当", 
         "「決済はPayPayとかLINE Payとかクレカとか今風の全部！手数料高いなら銀行振込だけでいいかも。あとポイントな！100円で何点か？普通でいいよ、後で変えられるようにしといて。」",
         "※決済代行の審査リードタイム無視。ポイントの有効期限や会計処理も白紙。")
    ]

    positions = [
        (Inches(0.8), Inches(1.5)),
        (Inches(6.86), Inches(1.5)),
        (Inches(0.8), Inches(4.2)),
        (Inches(6.86), Inches(4.2))
    ]

    for (title, quote, note), (gx, gy) in zip(hearings, positions):
        add_card(s3, gx, gy, grid_w, grid_h, bg_color=C_CARD_BG, border_color=C_BORDER)
        tb = s3.shapes.add_textbox(gx + Inches(0.25), gy + Inches(0.2), grid_w - Inches(0.5), grid_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_PRIMARY

        p = tf.add_paragraph()
        p.text = quote
        p.font.size = Pt(10.5)
        p.font.italic = True
        p.font.color.rgb = C_TEXT_MAIN

        p = tf.add_paragraph()
        p.text = note
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_ORANGE

    # ==========================================
    # SLIDE 4: システム要件一覧（曖昧仕様の整理）
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "システム要件定義書（要確認項目一覧）")

    req_table = [
        ("商品・在庫管理", "店頭の帳面と手動連動（？）", "店頭在庫とEC在庫は別。無くなったら店主が手動で管理画面から在庫0にする（忘れると無限に買える）"),
        ("カート・注文", "数量バリデーションなし", "注文数量の最小値・最大値制限未定（マイナス値入力の考慮漏れ）"),
        ("消費税計算", "端数処理ポリシー未統一", "画面によって切り捨て・四捨五入が混在しているが『1円くらい気にしない』とのこと"),
        ("会員・ポイント", "仕様未確定のまま突切り", "付与タイミング、失効期限、利用下限など全て『あとで決める』"),
        ("ギフト・熨斗", "備考欄に手書き対応", "熨斗選択UIは作らず、備考欄に『お中元』と書かれたら店主が手書きで熨斗を貼る運用"),
        ("セキュリティ", "予算ゼロのため最低限", "認証不備、XSS、生SQL等の脆弱性検査は工数外として除外")
    ]

    # テーブル作成
    rows = 7
    cols = 3
    t_left = Inches(0.8)
    t_top = Inches(1.5)
    t_w = Inches(11.733)
    t_h = Inches(5.2)

    table_shape = s4.shapes.add_table(rows, cols, t_left, t_top, t_w, t_h)
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(3.2)
    table.columns[2].width = Inches(6.333)

    headers = ["機能分類", "要件・仕様（ヒアリング結果）", "システム設計上の実態・懸念点"]
    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_WHITE

    for ri, (c1, c2, c3) in enumerate(req_table):
        for ci, val in enumerate([c1, c2, c3]):
            cell = table.cell(ri + 1, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_WHITE if ri % 2 == 0 else C_BG_LIGHT
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.color.rgb = C_TEXT_MAIN

    # ==========================================
    # SLIDE 5: 破綻したサイトマップ（見た目だけ整然）
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "ECサイト サイトマップ設計（構造破綻・孤立ページあり）")

    # サイトマップのボックス配置
    sm_nodes = [
        ("TOPページ\n(秘伝タレ動画音ドーン)", Inches(5.0), Inches(1.5), Inches(3.3), Inches(0.9), C_NAVY_PRIMARY, C_WHITE),
        
        ("商品一覧\n(/products)", Inches(1.0), Inches(2.8), Inches(2.5), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("店主の熱いこだわり\n(/kodawari)", Inches(3.8), Inches(2.8), Inches(2.5), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("店舗案内\n(Mapスクショ画像)", Inches(6.6), Inches(2.8), Inches(2.5), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("マイページ\n(IDOR脆弱性あり)", Inches(9.4), Inches(2.8), Inches(2.5), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),

        ("商品詳細\n(/products/{id})", Inches(1.0), Inches(4.0), Inches(2.5), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("買い物かご\n(マイナス数量バグ)", Inches(3.8), Inches(4.0), Inches(2.5), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("お届け先・決済入力\n(/checkout)", Inches(6.6), Inches(4.0), Inches(2.5), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),
        ("注文完了\n(/thanks)", Inches(9.4), Inches(4.0), Inches(2.5), Inches(0.8), C_CARD_BG, C_TEXT_MAIN),

        ("【孤立】常連専用隠し部屋\n(ヘッダー/フッターに導線なし)", Inches(1.0), Inches(5.4), Inches(4.0), Inches(1.2), C_ALERT_BG, C_ORANGE),
        ("【法的表記欠落】特定商取引法 / 利用規約\n(導線ゼロ・テンプレURL直打ちのみ閲覧可能)", Inches(5.4), Inches(5.4), Inches(6.5), Inches(1.2), C_ALERT_BG, C_ORANGE)
    ]

    for title, x, y, w, h, bg, tc in sm_nodes:
        shape = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg
        shape.line.color.rgb = C_ORANGE if bg == C_ALERT_BG else C_BORDER
        shape.line.width = Pt(1.5) if bg == C_ALERT_BG else Pt(1)
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = tc

    # ==========================================
    # SLIDE 6: 開発スケジュール（超突貫デスマーチ）
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "ローンチまでの推進スケジュール（超突貫工程）")

    sched_data = [
        ("Phase 1: 要件定義", "Day 1 (居酒屋)", "店主とビールを飲みながら口頭で仕様決定。議事録なし。"),
        ("Phase 2: 実装・コーディング", "Day 2 - Day 4", "1500行の神コントローラに全処理を殴り書き。動けばヨシ！"),
        ("Phase 3: テスト・品質管理", "Day 5 (1時間)", "店主が自分のPCで1回カートに入れたら購入できたのでテスト完了。"),
        ("Phase 4: 本番ローンチ", "Day 6 (大安)", "本番デバッグモード(APP_DEBUG=true)のまま奇跡の公開！")
    ]

    sw = Inches(2.65)
    sgap = Inches(0.35)
    for i, (phase, period, desc) in enumerate(sched_data):
        sx = Inches(0.8) + i * (sw + sgap)
        sy = Inches(1.8)
        sh = Inches(4.5)

        add_card(s6, sx, sy, sw, sh)
        tb = s6.shapes.add_textbox(sx + Inches(0.2), sy + Inches(0.2), sw - Inches(0.4), sh - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"STEP {i+1}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_ORANGE

        p = tf.add_paragraph()
        p.text = phase
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_DARK

        p = tf.add_paragraph()
        p.text = period
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_PRIMARY

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_TEXT_MAIN

    # 保存
    output_path = "/Users/gsgwr/.gemini/antigravity/scratch/gondawara-tsukudani/docs/planning/gondawara_dx_proposal.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
